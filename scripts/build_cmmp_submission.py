"""Convert the current, checked LaTeX manuscript to CMMP's editable Word format.

Requires pandoc, python-docx and a resolved pdfLaTeX .aux file.  Run with the
bundled document Python runtime.  Private editorial contact details are never
read by this converter.  Figures are collected at the end, as CMMP requests.
"""
from __future__ import annotations
import argparse
import hashlib
import json
from pathlib import Path
import re
import subprocess
import zipfile
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
MANUSCRIPT = ROOT / 'manuscript'

def braced(text, start):
    assert text[start] == '{'
    depth = 1
    pos = start + 1
    while depth:
        if text[pos] == '{' and (pos == 0 or text[pos - 1] != '\\'):
            depth += 1
        elif text[pos] == '}' and text[pos - 1] != '\\':
            depth -= 1
        pos += 1
    return text[start + 1:pos - 1], pos

def caption(block, prefix):
    found = re.search(r'\\caption\s*\{', block)
    assert found, block[:100]
    content, end = braced(block, found.end() - 1)
    return block[:found.start()] + '\\caption{' + prefix + ' ' + content.rstrip('.') + '}' + block[end:]

def reference_doc(path):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.left_margin = section.right_margin = Mm(20)
    section.top_margin = section.bottom_margin = Mm(20)
    for name in ('Normal', 'Body Text', 'First Paragraph'):
        style = doc.styles[name] if name in doc.styles else doc.styles['Normal']
        style.font.name = 'Times New Roman'
        style.font.size = Pt(10)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.line_spacing = 1.15
    for name, size in [('Title', 16), ('Heading 1', 12), ('Heading 2', 11), ('Heading 3', 10), ('Heading 4', 10), ('Heading 5', 10)]:
        style = doc.styles[name]
        style.font.name = 'Times New Roman'
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for name in ('Caption', 'Table Caption', 'Image Caption'):
        if name in doc.styles:
            doc.styles[name].font.name = 'Times New Roman'
            doc.styles[name].font.size = Pt(9)
    doc.save(path)

def prepare(source, aux):
    labels = dict(re.findall(r'\\newlabel\{([^}]+)\}\{\{([^}]+)\}', aux))
    cites = dict(re.findall(r'\\bibcite\{([^}]+)\}\{\{(\d+)\}', aux))
    authors = dict(re.findall(r'\\bibitem\[([^\]]+)\]\{([^}]+)\}', source))
    authors = {key: label.split('(')[0] for label, key in authors.items()}
    expanded = re.sub(r'\\input\{([^}]+)\}', lambda m: (MANUSCRIPT / m[1]).read_text(encoding='utf-8'), source)
    # Conversion operates on a private expanded copy; the scientific master stays intact.
    preamble, body = expanded.split('\\begin{document}', 1)
    body = body.rsplit('\\end{document}', 1)[0]
    body = re.sub(r'\\ifdefined\\BlindReview.*?\\else(.*?)\\fi', r'\1', body, flags=re.S)
    body = re.sub(r'\\(?:eqref|ref)\{([^}]+)\}', lambda m: ('(' + labels[m[1]] + ')') if m[0].startswith('\\eqref') else labels[m[1]], body)
    def citation(m):
        keys = m[2].split(',')
        numbers = '[' + ', '.join(cites[key] for key in keys) + ']'
        return (authors[keys[0]] + ' ' if m[1] == 'citet' else '') + numbers
    body = re.sub(r'\\(citep|citet|cite)\{([^}]+)\}', citation, body)
    figures = []
    def figure(m):
        block = m[0]
        label = re.search(r'\\label\{([^}]+)\}', block)[1]
        number = labels[label]
        block = caption(block, 'Fig. ' + number)
        block = re.sub(r'\\includegraphics(?:\[[^]]*\])?\{([^}]+)\}', lambda x: '\\includegraphics[width=17cm]{' + str(Path(x[1]).with_suffix('.png')).replace('\\', '/') + '}', block)
        figures.append(block)
        return ''
    body = re.sub(r'\\begin\{figure\}(?:\[[^]]*\])?.*?\\end\{figure\}', figure, body, flags=re.S)
    def table(m):
        block = m[0]
        label = re.search(r'\\label\{([^}]+)\}', block)[1]
        return caption(block, 'Table ' + labels[label])
    body = re.sub(r'\\begin\{table\}(?:\[[^]]*\])?.*?\\end\{table\}', table, body, flags=re.S)
    equation_no = 0
    def equation(m):
        nonlocal equation_no
        equation_no += 1
        expression = m[1]
        label = re.search(r'\\label\{([^}]+)\}', expression)
        number = labels[label[1]] if label else str(equation_no)
        return '\\[' + expression + '\\qquad\\text{(' + number + ')}\\]'
    body = re.sub(r'\\begin\{equation\}(.*?)\\end\{equation\}', equation, body, flags=re.S)
    prop_no = 0
    def proposition(m):
        nonlocal prop_no
        prop_no += 1
        return '\\paragraph*{Proposition ' + str(prop_no) + '. ' + m[1] + '}\n'
    body = re.sub(r'\\begin\{proposition\}\[([^]]+)\]', proposition, body)
    body = body.replace('\\end{proposition}', '')
    body = body.replace('\\begin{proof}', '\\paragraph*{Proof.}').replace('\\end{proof}', '\\hfill$\\square$')
    body = re.sub(r'\\begin\{thebibliography\}\{[^}]+\}', lambda _: '\\section*{References}', body)
    body = body.replace('\\end{thebibliography}', '')
    body = re.sub(r'\\bibitem(?:\[[^]]+\])?\{([^}]+)\}', lambda m: '\n\\noindent\\textbf{[' + cites[m[1]] + ']} ', body)
    # Give Word explicit section numbering, including A.1 etc., rather than fields.
    section_no = subsection_no = 0
    appendix = False
    def section(m):
        nonlocal section_no, subsection_no, appendix
        if m[0] == '\\appendix':
            appendix = True
            section_no = subsection_no = 0
            return ''
        level, star, title = m[1], m[2], m[3]
        if star:
            return m[0]
        if level == 'section':
            section_no += 1
            subsection_no = 0
            prefix = chr(64 + section_no) if appendix else str(section_no)
        else:
            subsection_no += 1
            prefix = (chr(64 + section_no) if appendix else str(section_no)) + '.' + str(subsection_no)
        return '\\' + level + '*{' + prefix + ' ' + title + '}'
    body = re.sub(r'\\appendix|\\(section|subsection)(\*?)\{([^}]+)\}', section, body)
    body = re.sub(r'\\label\{[^}]+\}', '', body)
    body = re.sub(r'\\(?:begin|end)\{samepage\}', '', body)
    body = body.replace('\\phantomsection', '').replace('\\maketitle', '')
    # TeX's legacy font switches are equivalent to explicit math text here.
    body = re.sub(r'\\rm\s+([A-Za-z]+)', lambda m: '\\mathrm{' + m[1] + '}', body)
    body = body.replace('\\hbox{', '\\text{')
    body += '\n\\clearpage\n\\section*{Figure captions and illustrations}\n' + '\n\\clearpage\n'.join(figures)
    preamble = '''\\documentclass{article}
\\usepackage{amsmath,amssymb,booktabs,array,graphicx,hyperref}
\\newcommand{\\dd}{\\,\\mathrm{d}}
\\newcommand{\\R}{\\mathbb{R}}
\\title{Geometry-consistent conservative transfer of potential-flow states to an embedded-boundary volume-of-fluid solver}
\\author{Saveliy Baturin\\\\Independent Researcher, Moscow, Russia\\\\saveliy.xo@gmail.com\\\\ORCID: 0009-0005-4410-9520}
\\date{}
'''
    return preamble + '\\begin{document}\n' + body + '\n\\end{document}\n', len(figures)

def polish(path):
    doc = Document(path)
    # Word's default Title style can carry an unrelated blue separator.
    for border in doc.styles.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    for border in doc.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    drawing_count = 0
    for p in doc.paragraphs:
        p.paragraph_format.widow_control = True
        if p.text.startswith('Table '):
            p.paragraph_format.keep_with_next = True
        if p.style.name == 'Title':
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.name = 'Times New Roman'
                run.font.size = Pt(16)
                run.font.color.rgb = RGBColor(0, 0, 0)
        if p.style.name in ('Author', 'Date'):
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if p.text == 'Abstract':
            p.style = doc.styles['Heading 2']
        if p.text == 'Figure captions and illustrations':
            p.paragraph_format.page_break_before = True
        if p._p.xpath('.//w:drawing'):
            if drawing_count:
                p.paragraph_format.page_break_before = True
            p.paragraph_format.keep_with_next = True
            drawing_count += 1
    for table in doc.tables:
        table.autofit = True
        for row_no, row in enumerate(table.rows):
            trpr = row._tr.get_or_add_trPr()
            trpr.append(OxmlElement('w:cantSplit'))
            if row_no == 0:
                trpr.append(OxmlElement('w:tblHeader'))
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.line_spacing = 1.0
                    p.paragraph_format.keep_with_next = row_no < len(table.rows) - 1
                    for run in p.runs:
                        run.font.name = 'Times New Roman'
                        run.font.size = Pt(9)
    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    reference_numbers = [int(m[1]) for p in doc.paragraphs if (m := re.match(r'^\[(\d+)\]', p.text))]
    assert reference_numbers == list(range(1, 18)), reference_numbers
    doc.save(path)
    with zipfile.ZipFile(path) as z:
        xml = z.read('word/document.xml').decode()
        assert '<m:oMath' in xml, 'No editable native Word equations'
        assert '\\begin{' not in xml and '\\cite' not in xml, 'Raw LaTeX remains'
        return {'native_math_elements': xml.count('<m:oMath>'), 'tables': len(doc.tables), 'paragraphs': len(doc.paragraphs)}

def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--pandoc', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    out = args.output.resolve()
    out.mkdir(parents=True, exist_ok=True)
    source = (MANUSCRIPT / 'transfer_article.tex').read_text(encoding='utf-8')
    aux = (ROOT / 'build/transfer_article_pdf/transfer_article.aux').read_text(encoding='utf-8')
    prepared, figures = prepare(source, aux)
    work = ROOT / 'build/cmmp_word_conversion'
    work.mkdir(parents=True, exist_ok=True)
    reference = work / 'reference.docx'
    reference_doc(reference)
    expanded = work / 'expanded.tex'
    expanded.write_text(prepared, encoding='utf-8')
    target = out / 'transfer_article.docx'
    command = [str(args.pandoc.resolve()), str(expanded), '--from=latex', '--to=docx', '--standalone',
               '--resource-path=' + str(MANUSCRIPT), '--reference-doc=' + str(reference), '--output=' + str(target)]
    run = subprocess.run(command, cwd=MANUSCRIPT, capture_output=True, text=True, encoding='utf-8')
    (work / 'pandoc.log').write_text(run.stdout + run.stderr, encoding='utf-8')
    if run.returncode:
        raise RuntimeError(run.stderr)
    if 'Could not convert TeX math' in run.stderr:
        raise RuntimeError('A mathematical expression failed native Word conversion: ' + run.stderr)
    result = polish(target)
    result.update(figures=figures, source_sha256=hashlib.sha256((MANUSCRIPT / 'transfer_article.tex').read_bytes()).hexdigest(),
                  docx_sha256=hashlib.sha256(target.read_bytes()).hexdigest(), pandoc_log=run.stderr)
    (work / 'conversion.json').write_text(json.dumps(result, indent=2) + '\n', encoding='utf-8')
    print(json.dumps(result, indent=2))

if __name__ == '__main__':
    main()
